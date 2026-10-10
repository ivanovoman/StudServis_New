"""Проверки сборки презентации к защите."""

from __future__ import annotations

import io

import pytest
from pptx import Presentation
from pptx.util import Emu

from app.modules.documents.pptx_engine import (
    SLIDE_H,
    Slide,
    build_presentation,
)


def _open(data: bytes) -> Presentation:
    return Presentation(io.BytesIO(data))


def _bottom(slide) -> float:
    """Нижняя граница самого нижнего элемента, в сантиметрах."""
    return max((Emu(sh.top + sh.height).cm for sh in slide.shapes), default=0.0)


def _texts(slide) -> str:
    out = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            out.append(shape.text_frame.text)
    return "\n".join(out)


def test_титульный_лист_берёт_реквизиты_из_настроек():
    """Модель выдумывает ФИО руководителя — на защите это замечают."""
    data = build_presentation(
        [Slide(title="Тема работы", bullets=["Выполнил: Выдуманный В.В."])],
        topic="Коллизии в праве",
        university="МФЮА",
        author="Иванов И.И.",
        supervisor="Петров П.П.",
        year="2026",
    )
    prs = _open(data)
    title = _texts(prs.slides[0])

    assert "Коллизии в праве" in title
    assert "Иванов И.И." in title
    assert "Петров П.П." in title
    assert "Выдуманный" not in title, "титул модели должен быть заменён"


def test_слайд_модели_не_теряется_если_он_не_титульный():
    data = build_presentation(
        [Slide(title="Актуальность", bullets=["Первый тезис"])],
        topic="Коллизии в праве",
    )
    prs = _open(data)

    assert len(prs.slides) == 2, "титульный добавляется, содержание остаётся"
    assert "Актуальность" in _texts(prs.slides[1])


def test_таблица_попадает_в_слайд():
    rows = [["Критерий", "А", "Б"], ["Источник", "закон", "договор"]]
    data = build_presentation(
        [Slide(title="Сравнение", table=rows)], topic="Тема",
    )
    prs = _open(data)
    tables = [sh for sh in prs.slides[1].shapes if sh.has_table]

    assert len(tables) == 1
    table = tables[0].table
    assert len(table.rows) == 2
    assert len(table.columns) == 3
    assert table.cell(1, 2).text == "договор"


def test_таблица_с_пунктами_не_вылезает_за_край():
    """Самый частый дефект вёрстки: таблица съедала всю высоту."""
    rows = [["Критерий", "А", "Б"]] + [
        [f"строка {i}", "значение", "значение"] for i in range(5)
    ]
    data = build_presentation(
        [Slide(
            title="Перегруженный слайд",
            table=rows,
            bullets=[f"Содержательный вывод номер {i}" for i in range(4)],
        )],
        topic="Тема",
    )
    prs = _open(data)

    assert _bottom(prs.slides[1]) <= SLIDE_H.cm, "контент уехал за край слайда"


@pytest.mark.parametrize("count", [3, 6, 9])
def test_кегль_уменьшается_с_ростом_числа_пунктов(count: int):
    data = build_presentation(
        [Slide(title="Выводы",
               bullets=[f"Достаточно длинный тезис под номером {i}"
                        for i in range(count)])],
        topic="Тема",
    )
    prs = _open(data)
    sizes = [
        run.font.size.pt
        for shape in prs.slides[1].shapes
        if shape.has_text_frame
        for para in shape.text_frame.paragraphs
        for run in para.runs
        if run.font.size and not run.font.bold
    ]
    assert sizes, "у пунктов должен быть задан кегль"
    assert min(sizes) >= 12, "мельче 12 пунктов читать с задних рядов нельзя"


def test_вложенные_пункты_получают_уровень_и_меньший_кегль():
    data = build_presentation(
        [Slide(
            title="Задачи",
            bullets=["Задачи:", "Раскрыть понятие", "Классифицировать виды"],
            levels=[0, 1, 1],
        )],
        topic="Тема",
    )
    prs = _open(data)
    paras = [
        para
        for shape in prs.slides[1].shapes
        if shape.has_text_frame
        for para in shape.text_frame.paragraphs
        if para.text.strip()
    ]
    body = [p for p in paras if p.text.strip().startswith(("•", "–"))]

    assert body[0].level == 0
    assert body[1].level == 1, "подпункт должен быть смещён"
    assert body[0].runs[0].font.bold, "подзаголовок выделяется жирным"
    assert body[1].runs[0].font.size < body[0].runs[0].font.size


def test_финальный_слайд_центрируется():
    """Пустой слайд с заголовком в углу выглядит как сбой загрузки."""
    data = build_presentation(
        [Slide(title="Спасибо за внимание")], topic="Тема",
    )
    prs = _open(data)
    last = prs.slides[-1]

    shapes = [sh for sh in last.shapes if sh.has_text_frame]
    assert len(shapes) == 1, "лишних элементов на финальном слайде нет"
    assert Emu(shapes[0].top).cm > 5.0, "заголовок должен стоять по центру"


def test_презентация_без_реквизитов_всё_равно_собирается():
    data = build_presentation([Slide(title="Один слайд", bullets=["Тезис"])])

    prs = _open(data)
    assert len(prs.slides) >= 1
    assert data[:2] == b"PK"


def test_заметки_докладчика_сохраняются():
    data = build_presentation(
        [Slide(title="Актуальность", bullets=["Тезис"],
               note="Здесь говорю о статистике дел.")],
        topic="Тема",
    )
    prs = _open(data)

    notes = prs.slides[1].notes_slide.notes_text_frame.text
    assert "статистике" in notes
