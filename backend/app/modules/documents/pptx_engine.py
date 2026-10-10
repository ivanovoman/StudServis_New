"""Сборка презентации к защите в формате .pptx.

Шаблоны PowerPoint по умолчанию рассчитаны на деловую презентацию с
картинками. Защита курсовой — другой жанр: проектор в аудитории часто
тусклый, задние ряды далеко, а комиссия смотрит на экран мельком,
между чтением самой работы. Поэтому здесь крупный шрифт, тёмный текст
на белом и никаких декоративных элементов: всё, что не несёт смысла,
на плохом проекторе превращается в грязь.

Размер слайда — 16:9. Аудиторные проекторы бывают и 4:3, но
широкоэкранная презентация на них показывается с полями, тогда как
4:3 на широком экране растягивается и текст плывёт.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Cm, Pt

# Размеры слайда 16:9 в сантиметрах.
SLIDE_W = Cm(33.87)
SLIDE_H = Cm(19.05)

MARGIN_X = Cm(2.0)
TITLE_TOP = Cm(1.2)
BODY_TOP = Cm(4.2)

TEXT_COLOR = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
MUTED = RGBColor(0x55, 0x55, 0x55)
LINE = RGBColor(0xC8, 0xC8, 0xC8)

FONT = "Times New Roman"


@dataclass
class Slide:
    title: str = ""
    bullets: list[str] = field(default_factory=list)
    table: list[list[str]] = field(default_factory=list)
    note: str = ""
    # Уровень вложенности каждого пункта: 0 — обычный, 1 — подпункт.
    # Без него «Задачи:» и сами задачи выглядели одним плоским
    # списком, где заголовок неотличим от содержания.
    levels: list[int] = field(default_factory=list)


def _blank(prs: Presentation):
    """Пустой макет: готовые плейсхолдеры навязывают свои отступы."""
    return prs.slides.add_slide(prs.slide_layouts[6])


# Запас по высоте: оценка длины текста приблизительна, и упереться
# ровно в край опаснее, чем оставить поле.
MIN_BODY_PT = 12
SPACE_AFTER_PT = 14


def _needed_height_cm(count: int, size_pt: float, width_cm: float,
                      longest: int) -> float:
    """Сколько места займут пункты при данном кегле.

    Длинный пункт переносится на несколько строк, и считать его за
    одну — главный источник переполнения. Ширину символа берём как
    0.5 кегля: для Times New Roman на кириллице это близко к правде.
    """
    char_w_cm = size_pt * 0.5 / 28.35
    per_line = max(1, int(width_cm / char_w_cm))
    lines = 0
    for _ in range(count):
        lines += max(1, -(-longest // per_line))
    return lines * size_pt * 1.15 / 28.35 + count * SPACE_AFTER_PT / 28.35


def _body_font_size(slide: Slide, available_cm: float | None = None) -> Pt:
    """Кегль под объём текста и оставшееся на слайде место.

    Фиксированный размер приводит к одному из двух: либо короткий слайд
    выглядит пустым, либо длинный вылезает за край.

    Отдельная беда — слайд с таблицей: она съедает больше половины
    высоты, и пункты под ней переставали помещаться, уезжая за нижний
    край. Поэтому кегль подбирается от реально свободного места, а не
    только от числа пунктов.
    """
    count = len(slide.bullets)
    if not count:
        return Pt(24)
    longest = max(len(b) for b in slide.bullets)

    if count <= 3 and longest <= 70:
        size = 28
    elif count <= 5 and longest <= 110:
        size = 24
    elif count <= 6:
        size = 20
    else:
        size = 18

    if available_cm is None:
        return Pt(size)

    width_cm = SLIDE_W.cm - 2 * MARGIN_X.cm - 1.0
    while size > MIN_BODY_PT and _needed_height_cm(
            count, size, width_cm, longest) > available_cm:
        size -= 1
    return Pt(size)


def _add_title(slide_obj, text: str, *, big: bool = False) -> None:
    box = slide_obj.shapes.add_textbox(
        MARGIN_X, TITLE_TOP, SLIDE_W - 2 * MARGIN_X,
        Cm(5.0) if big else Cm(2.6),
    )
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = frame.paragraphs[0]
    para.text = text
    para.alignment = PP_ALIGN.CENTER if big else PP_ALIGN.LEFT
    run = para.runs[0]
    run.font.name = FONT
    run.font.bold = True
    run.font.color.rgb = ACCENT
    # Длинный заголовок уменьшаем: перенос на три строки съедает место
    # под содержание и выглядит как ошибка вёрстки.
    if big:
        run.font.size = Pt(40) if len(text) <= 80 else Pt(32)
    else:
        run.font.size = Pt(32) if len(text) <= 55 else Pt(26)


def _add_rule(slide_obj) -> None:
    """Линия под заголовком — единственный декоративный элемент.

    Нужна не для красоты: она отделяет заголовок от пунктов, когда
    проектор заваливает контраст и цвет заголовка перестаёт читаться
    как отдельный уровень.
    """
    from pptx.enum.shapes import MSO_SHAPE

    line = slide_obj.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, MARGIN_X, Cm(3.6),
        SLIDE_W - 2 * MARGIN_X, Pt(1.5),
    )
    line.fill.solid()
    line.fill.fore_color.rgb = LINE
    line.line.fill.background()
    line.shadow.inherit = False


def _add_bullets(slide_obj, slide: Slide, top: Cm) -> None:
    available_cm = SLIDE_H.cm - top.cm - 1.5
    box = slide_obj.shapes.add_textbox(
        MARGIN_X, top, SLIDE_W - 2 * MARGIN_X, Cm(available_cm),
    )
    frame = box.text_frame
    frame.word_wrap = True
    size = _body_font_size(slide, available_cm)

    levels = slide.levels or [0] * len(slide.bullets)

    for i, text in enumerate(slide.bullets):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        level = levels[i] if i < len(levels) else 0
        # Маркер ставим символом, а не списком: настоящая нумерация в
        # pptx требует правки XML, а визуально разницы нет.
        marker = "–" if level else "•"
        para.text = f"{marker}  {text}"
        para.level = min(level, 4)
        # Подпункты теснее: они продолжают мысль строки над собой, и
        # полный отбив рвал бы связь между ними.
        para.space_after = Pt(8) if level else Pt(14)
        para.line_spacing = 1.15
        run = para.runs[0]
        run.font.name = FONT
        run.font.size = Pt(size.pt - 3) if level else size
        run.font.color.rgb = TEXT_COLOR
        # Подзаголовок (строка с двоеточием на конце) выделяем: без
        # этого он теряется среди своих же подпунктов.
        if not level and text.rstrip().endswith(":"):
            run.font.bold = True


def _add_table(slide_obj, rows: list[list[str]], top: Cm,
               reserve_cm: float = 0.0) -> Cm:
    """Рисует таблицу и возвращает отметку, с которой можно писать дальше.

    reserve_cm — сколько высоты надо оставить под пункты ниже. Без
    этого таблица на шесть строк занимала всё место, и текст под ней
    уходил за край слайда.
    """
    n_rows, n_cols = len(rows), len(rows[0])
    limit = SLIDE_H.cm - top.cm - 1.5 - reserve_cm
    height = Cm(max(3.0, min(1.4 * n_rows, 10.0, limit)))
    shape = slide_obj.shapes.add_table(
        n_rows, n_cols, MARGIN_X, top, SLIDE_W - 2 * MARGIN_X, height,
    )
    table = shape.table

    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            para = cell.text_frame.paragraphs[0]
            para.alignment = PP_ALIGN.LEFT
            if not para.runs:
                continue
            run = para.runs[0]
            run.font.name = FONT
            # Текст в клетках мельче пунктов: таблица и так читается
            # как единый блок, а крупный шрифт заставит её вылезти.
            run.font.size = Pt(18) if n_rows <= 5 else Pt(14)
            run.font.bold = r == 0
            run.font.color.rgb = TEXT_COLOR

    return Cm(top.cm + height.cm + 0.8)


def _add_note(slide_obj, text: str) -> None:
    if not text:
        return
    frame = slide_obj.notes_slide.notes_text_frame
    frame.text = text


def build_presentation(
    slides: list[Slide],
    *,
    topic: str = "",
    university: str = "",
    author: str = "",
    supervisor: str = "",
    year: str = "",
) -> bytes:
    """Собирает .pptx и возвращает его байтами.

    Титульный слайд формируется из переданных реквизитов, а не из
    текста модели: модель склонна выдумывать ФИО руководителя, а это
    тот редкий случай, когда выдумка заметна комиссии мгновенно.
    """
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    body = list(slides)

    # Титульный лист: если первым слайдом модель уже дала титульный,
    # заменяем его своим — с настоящими реквизитами.
    first_is_title = bool(body) and _looks_like_title(body[0], topic)
    if first_is_title:
        model_title = body.pop(0)
    else:
        model_title = None

    title_slide = _blank(prs)
    _add_title(title_slide, topic or (model_title.title if model_title else "Презентация"), big=True)

    lines = []
    if university:
        lines.append(university)
    if author:
        lines.append(f"Выполнил: {author}")
    if supervisor:
        lines.append(f"Научный руководитель: {supervisor}")
    if year:
        lines.append(str(year))

    if lines:
        box = title_slide.shapes.add_textbox(
            MARGIN_X, Cm(10.5), SLIDE_W - 2 * MARGIN_X, Cm(6.0),
        )
        frame = box.text_frame
        frame.word_wrap = True
        for i, text in enumerate(lines):
            para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            para.text = text
            para.alignment = PP_ALIGN.CENTER
            para.space_after = Pt(10)
            run = para.runs[0]
            run.font.name = FONT
            run.font.size = Pt(20)
            run.font.color.rgb = MUTED

    for slide in body:
        obj = _blank(prs)

        # Слайд без содержания — финальное «Спасибо за внимание».
        # С обычной вёрсткой заголовок повисал в левом верхнем углу
        # пустого экрана и выглядел как недогрузившийся слайд.
        if slide.title and not slide.bullets and not slide.table:
            box = obj.shapes.add_textbox(
                MARGIN_X, Cm(7.5), SLIDE_W - 2 * MARGIN_X, Cm(4.0))
            frame = box.text_frame
            frame.word_wrap = True
            frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            para = frame.paragraphs[0]
            para.text = slide.title
            para.alignment = PP_ALIGN.CENTER
            run = para.runs[0]
            run.font.name = FONT
            run.font.size = Pt(40)
            run.font.bold = True
            run.font.color.rgb = ACCENT
            _add_note(obj, slide.note)
            continue

        if slide.title:
            _add_title(obj, slide.title)
            _add_rule(obj)
        top = BODY_TOP if slide.title else Cm(2.0)
        if slide.table:
            # Под пункты резервируем место заранее: иначе таблица
            # заберёт всю высоту и текст под ней окажется за кадром.
            reserve = 0.0
            if slide.bullets:
                reserve = min(6.0, 1.2 * len(slide.bullets) + 0.8)
            top = _add_table(obj, slide.table, top, reserve)
        if slide.bullets:
            _add_bullets(obj, slide, top)
        _add_note(obj, slide.note)

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def _looks_like_title(slide: Slide, topic: str) -> bool:
    """Похож ли первый слайд на титульный.

    Признак — отсутствие содержательных пунктов: на титульном стоят
    реквизиты («Выполнил», «Руководитель»), а не тезисы.
    """
    if slide.table:
        return False
    marks = ("выполнил", "руководитель", "студент", "курсовая", "работа")
    joined = " ".join(slide.bullets).lower()
    if any(m in joined for m in marks):
        return True
    if topic and slide.title and topic.lower()[:40] in slide.title.lower():
        return True
    return False
